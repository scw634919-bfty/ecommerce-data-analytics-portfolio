"""Draw a static MOCKUP of the planned Power BI margin dashboard.

It reads the numbers from the recalculated Excel model, so the mockup shows the
project's real (simulated) results. It is a picture of the layout, not a Power BI file.

Usage: python3 make_dashboard_mockup.py   ->  images/dashboard_mockup.png
"""
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from openpyxl import load_workbook

HERE = Path(__file__).parent
wb = load_workbook(HERE / "B2B_Pricing_Margin_Model.xlsx", data_only=True)
ml, asr = wb["Margin_Leakage"], wb["Assortment"]

# ---------------------------------------------------------------- palette (validated)
SURFACE = "#fcfcfb"
PAGE = "#f0efec"
INK = "#0b0b0b"
INK2 = "#52514e"
MUTED = "#8a8984"
GRID = "#e4e3df"
BLUE = "#2a78d6"     # categorical slot 1: actual / totals / positive
ORANGE = "#eb6834"   # categorical slot 2: decreases
RED = "#e34948"      # diverging negative pole
plt.rcParams.update({"font.family": "Arial", "font.size": 10, "text.color": INK,
                     "axes.labelcolor": INK2, "xtick.color": INK2, "ytick.color": INK2})


def find_rows(label_col_first, n):
    return range(label_col_first, label_col_first + n)


def section_start(title):
    for r in range(1, ml.max_row + 1):
        if ml[f"A{r}"].value == title:
            return r + 2
    raise ValueError(title)


# ---------------------------------------------------------------- data
kpi = {ml[f"E{r}"].value: ml[f"F{r}"].value for r in range(5, 13)}
bridge = [("GP at list", ml["B6"].value, "total"),
          ("Discounts", ml["B7"].value, "delta"),
          ("Cost drift", ml["B9"].value, "delta"),
          ("FX", ml["B10"].value, "delta"),
          ("Actual GP", ml["B11"].value, "total")]
r0 = section_start("2. By region")
regions = [(ml[f"A{r}"].value, ml[f"F{r}"].value) for r in find_rows(r0, 5)]
r0 = section_start("4. By sales rep")
reps = [(ml[f"A{r}"].value, ml[f"D{r}"].value, ml[f"E{r}"].value, ml[f"L{r}"].value) for r in find_rows(r0, 6)]
r0 = section_start("5. By month")
months = [(ml[f"A{r}"].value, ml[f"D{r}"].value, ml[f"E{r}"].value) for r in find_rows(r0, 12)]
flagged = []
for r in range(5, asr.max_row + 1):
    act = asr[f"O{r}"].value
    if act and not act.startswith(("Keep", "Preferred")):
        flagged.append((asr[f"A{r}"].value, asr[f"B{r}"].value, asr[f"I{r}"].value, act.split(":")[0]))

# ---------------------------------------------------------------- figure
fig = plt.figure(figsize=(16, 10), dpi=150, facecolor=PAGE)


def card(x, y, w, h):
    fig.add_artist(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0,rounding_size=0.006",
                                  transform=fig.transFigure, facecolor=SURFACE, edgecolor=GRID, lw=1, zorder=-10))


def style(ax, title):
    ax.set_facecolor(SURFACE)
    for s in ["top", "right", "left"]:
        ax.spines[s].set_visible(False)
    ax.spines["bottom"].set_color(GRID)
    ax.tick_params(length=0)
    ax.grid(axis="y", color=GRID, lw=0.8)
    ax.set_axisbelow(True)
    ax.set_title(title, loc="left", fontsize=12, fontweight="bold", color=INK, pad=10)


# header
fig.text(0.02, 0.955, "B2B Margin Dashboard", fontsize=20, fontweight="bold", color=INK)
fig.text(0.02, 0.928, "MOCKUP of the planned Power BI report  ·  simulated data, Oct 2025 – Sep 2026",
         fontsize=10.5, color=INK2)
for i, lab in enumerate(["Region: All", "Category: All", "Sales rep: All", "Month: All"]):
    x = 0.60 + i * 0.097
    fig.add_artist(FancyBboxPatch((x, 0.93), 0.088, 0.035, boxstyle="round,pad=0,rounding_size=0.004",
                                  transform=fig.transFigure, facecolor=SURFACE, edgecolor="#c9c8c2", lw=1, zorder=-10))
    fig.text(x + 0.008, 0.9475, lab + "  ▼", fontsize=9.5, color=INK2, va="center")

# KPI tiles
tiles = [("Revenue", f"${kpi['Revenue (USD)']/1e6:,.1f}M", "12 months, 5 regions"),
         ("Gross margin", f"{kpi['Gross margin %']*100:.1f}%",
          f"Target {kpi['Target margin % (revenue-weighted)']*100:.1f}%  ·  "
          f"{(kpi['Gross margin %']-kpi['Target margin % (revenue-weighted)'])*100:+.1f} pts"),
         ("Shortfall vs target", f"${kpi['Shortfall vs target (USD)']/1e3:,.0f}K", "Gross profit missed vs policy"),
         ("Below-floor lines, no approval", f"{kpi['  of which no approval logged']:.0f}",
          f"of {kpi['Order lines below floor']:.0f} below-floor order lines")]
for i, (t, v, s) in enumerate(tiles):
    x = 0.02 + i * 0.2425
    card(x, 0.775, 0.2325, 0.125)
    fig.text(x + 0.012, 0.872, t, fontsize=10.5, color=INK2)
    fig.text(x + 0.012, 0.815, v, fontsize=26, fontweight="bold", color=INK)
    fig.text(x + 0.012, 0.788, s, fontsize=9.5, color=MUTED)

# 1. margin bridge (waterfall)
card(0.02, 0.405, 0.475, 0.345)
ax = fig.add_axes([0.065, 0.44, 0.415, 0.255])
style(ax, "Where margin went: list price → actual gross profit ($K)")
run = 0
for i, (lab, val, kind) in enumerate(bridge):
    if kind == "total":
        bottom, height, col = 0, val, BLUE
        run = val
    else:
        bottom, height, col = run + val, -val, ORANGE
        run += val
    ax.bar(i, height / 1e3, bottom=bottom / 1e3, width=0.55, color=col, edgecolor=SURFACE, lw=2)
    ytxt = (bottom + height) / 1e3
    ax.text(i, ytxt + 70, f"{val/1e3:,.0f}" if kind == "total" else f"−{-val/1e3:,.0f}",
            ha="center", va="bottom", fontsize=10, color=INK)
ax.set_xticks(range(len(bridge)), [b[0] for b in bridge])
ax.set_ylim(0, 5600)
ax.yaxis.set_major_formatter(lambda v, _: f"{v:,.0f}")
fig.text(0.365, 0.711, "■", color=BLUE, fontsize=11)
fig.text(0.377, 0.711, "Gross profit", color=INK2, fontsize=9.5)
fig.text(0.432, 0.711, "■", color=ORANGE, fontsize=11)
fig.text(0.444, 0.711, "Margin lost", color=INK2, fontsize=9.5)

# 2. margin by month vs target
card(0.505, 0.405, 0.475, 0.345)
ax = fig.add_axes([0.545, 0.445, 0.42, 0.255])
style(ax, "Gross margin % by month vs target")
x = range(len(months))
ax.plot(x, [m[2] * 100 for m in months], color=MUTED, lw=2, ls=(0, (4, 3)))
ax.plot(x, [m[1] * 100 for m in months], color=BLUE, lw=2, marker="o", ms=5,
        markeredgecolor=SURFACE, markeredgewidth=1.5)
ax.text(11.15, months[-1][1] * 100, "Actual", color=INK, fontsize=9.5, va="center")
ax.text(11.15, months[-1][2] * 100 + 0.2, "Target", color=INK2, fontsize=9.5, va="center")
ax.set_xticks(list(x), [m[0][2:].replace("-", "/") for m in months], fontsize=9)
ax.set_xlim(-0.4, 12.2)
ax.set_ylim(9, 16)
ax.yaxis.set_major_formatter(lambda v, _: f"{v:.0f}%")
ax.annotate("Vendor cost +6% on 2 SKUs\n(price book not updated)", xy=(6, months[6][1] * 100),
            xytext=(3.2, 10.0), fontsize=9, color=INK2,
            arrowprops=dict(arrowstyle="-", color=MUTED, lw=1))

# 3. margin gap by region
card(0.02, 0.03, 0.31, 0.35)
ax = fig.add_axes([0.06, 0.095, 0.255, 0.215])
style(ax, "Margin vs target by region (pts)")
regions_sorted = sorted(regions, key=lambda r: r[1])
ax.barh([r[0] for r in regions_sorted], [r[1] for r in regions_sorted], height=0.55,
        color=[RED if r[1] < 0 else BLUE for r in regions_sorted], edgecolor=SURFACE, lw=2)
for i, (reg, gap) in enumerate(regions_sorted):
    ax.text(gap + (0.12 if gap >= 0 else -0.12), i, f"{gap:+.1f}", va="center",
            ha="left" if gap >= 0 else "right", fontsize=9.5, color=INK)
ax.axvline(0, color=INK2, lw=1)
ax.grid(axis="y", visible=False)
ax.grid(axis="x", color=GRID, lw=0.8)
ax.set_xlim(-3, 3.5)
fig.text(0.06, 0.042, "India: freight under-quoted (\$75 actual vs \$55 in price book)", fontsize=9, color=INK2)

# 4. reps
card(0.34, 0.03, 0.31, 0.35)
ax = fig.add_axes([0.375, 0.095, 0.26, 0.215])
style(ax, "Margin % by sales rep vs target")
for i, (rep, m, t, unap) in enumerate(reps):
    ax.bar(i, m * 100, width=0.55, color=BLUE, edgecolor=SURFACE, lw=2)
    ax.plot([i - 0.36, i + 0.36], [t * 100, t * 100], color=INK, lw=2)
    ax.text(i, m * 100 + 0.25, f"{m*100:.1f}%", ha="center", fontsize=9, color=INK)
ax.set_xticks(range(len(reps)), [r[0] for r in reps])
ax.set_ylim(8, 17.5)
ax.yaxis.set_major_formatter(lambda v, _: f"{v:.0f}%")
r04 = next(r for r in reps if r[0] == "R04")
fig.text(0.375, 0.042, f"—  target    ·    R04: {r04[3]:.0f} below-floor lines with no approval",
         fontsize=9, color=INK2)

# 5. flagged SKUs table
card(0.66, 0.03, 0.32, 0.35)
fig.text(0.672, 0.35, "SKUs to act on (assortment)", fontsize=12, fontweight="bold", color=INK)
cols_x = [0.672, 0.705, 0.865, 0.905]
for cx, h in zip(cols_x, ["SKU", "Model", "Margin", "Action"]):
    fig.text(cx, 0.318, h, fontsize=9.5, fontweight="bold", color=INK2)
fig.add_artist(plt.Line2D([0.672, 0.97], [0.311, 0.311], transform=fig.transFigure, color=GRID, lw=1))
icon = {"Below target": "▲", "Phase out": "■", "Consolidate": "●"}
for i, (sku, model, m, act) in enumerate(flagged[:8]):
    y = 0.288 - i * 0.031
    short = (model[:24] + "…") if len(model) > 25 else model
    fig.text(cols_x[0], y, sku, fontsize=9.5, color=INK)
    fig.text(cols_x[1], y, short, fontsize=9.5, color=INK)
    fig.text(cols_x[2], y, f"{m*100:.1f}%", fontsize=9.5, color=INK)
    fig.text(cols_x[3], y, f"{icon.get(act, '')} {act}", fontsize=9.5, color=INK)

fig.text(0.98, 0.008, "Mockup image generated from B2B_Pricing_Margin_Model.xlsx (simulated data)",
         ha="right", fontsize=8.5, color=MUTED)

out = HERE / "images"
out.mkdir(exist_ok=True)
fig.savefig(out / "dashboard_mockup.png", facecolor=PAGE)
print("saved", out / "dashboard_mockup.png", "| flagged:", len(flagged))
