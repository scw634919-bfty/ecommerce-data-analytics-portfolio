import sys, math
import pandas as pd
from openpyxl import load_workbook
F = sys.argv[1]
wb = load_workbook(F, data_only=True)
# 1) errors / empty formula results
wf = load_workbook(F)
errs, nones, nform = [], 0, 0
for ws in wf.worksheets:
    wv = wb[ws.title]
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                nform += 1
                v = wv[c.coordinate].value
                if v is None: nones += 1
                elif isinstance(v, str) and v.startswith("#"): errs.append((ws.title, c.coordinate, v))
print("formulas", nform, "errors", len(errs), errs[:10], "empty", nones)

# 2) independent pandas calc
D="data"
o=pd.read_csv(f"{D}/orders.csv"); sk=pd.read_csv(f"{D}/sku_master.csv"); cat=pd.read_csv(f"{D}/categories.csv")
o=o.merge(sk[["sku","category"]],on="sku").merge(cat,on="category")
o["rev"]=o.qty*o.quoted_price_local/o.fx_actual
o["rev_book"]=o.qty*o.quoted_price_local/o.fx_book
o["rev_list"]=o.qty*o.list_price_local/o.fx_book
o["cogs"]=o.qty*o.actual_landed_cost_usd; o["bcogs"]=o.qty*o.book_landed_cost_usd
o["gp"]=o.rev-o.cogs; o["m"]=o.gp/o.rev
o["below"]=(o.m<o.floor_margin).astype(int); o["unappr"]=((o.below==1)&(o.approval_logged=="N")).astype(int)
o["short"]=(o.rev*o.target_margin-o.gp).clip(lower=0)
ml=wb["Margin_Leakage"]
exp={"F5":o.rev.sum(),"F6":o.gp.sum(),"F7":o.gp.sum()/o.rev.sum(),"F9":o.short.sum(),"F10":o.below.sum(),"F11":o.unappr.sum(),
     "B6":o.rev_list.sum()-o.bcogs.sum(),"B7":(o.rev_book-o.rev_list).sum(),"B8":(o.rev_book-o.rev_list)[o.unappr==1].sum(),
     "B9":(o.bcogs-o.cogs).sum(),"B10":(o.rev-o.rev_book).sum(),"B11":o.gp.sum(),"B12":0}
bad=0
for k,v in exp.items():
    x=ml[k].value
    ok=abs(x-v)<=max(0.01,abs(v)*1e-9)
    bad+= not ok
    print(f"{k:4} excel={x:>16,.4f} pandas={v:>16,.4f} {'OK' if ok else 'MISMATCH'}")
# region table
for r in range(17,22):
    reg=ml[f"A{r}"].value; d=o[o.region==reg]
    for col,val in [("B",d.rev.sum()),("D",d.gp.sum()/d.rev.sum()),("G",d.short.sum())]:
        if abs(ml[f"{col}{r}"].value-val)>1e-6*max(1,abs(val)): bad+=1; print("region mismatch",reg,col)
# price book vs generator CSV
pb=pd.read_csv(f"{D}/price_book.csv"); ws=wb["Price_Book"]
for r in range(5,155):
    key=(ws[f"A{r}"].value,ws[f"D{r}"].value); row=pb[(pb.sku==key[0])&(pb.region==key[1])].iloc[0]
    if abs(ws[f"R{r}"].value-row.list_price_local)>0.001 or abs(ws[f"L{r}"].value-row.book_landed_cost_usd)>0.006:
        bad+=1; print("pricebook mismatch",key, ws[f"R{r}"].value,row.list_price_local)
print("MISMATCHES:",bad)
